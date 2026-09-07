"""Authenticated, expiring access to applicant document files."""

import tempfile
from io import BytesIO

from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.test import RequestFactory, TestCase, override_settings
from django.urls import reverse
from PIL import Image
from rest_framework.test import APIClient

from account.document_file_access import document_view_endpoint
from account.document_storage import (
    private_document_storage,
    remote_storage_configured,
    reset_storage_cache,
)
from account.models import (
    ApplicantDocument,
    ApplicantProfile,
    CustomUser,
    DocumentType,
    UserRole,
)
from account.services.signed_media import (
    DEFAULT_SIGNED_URL_TTL,
    signed_media_url,
    signed_url_ttl,
)


def _tiny_jpeg(name="ktp.jpg") -> ContentFile:
    buf = BytesIO()
    Image.new("RGB", (40, 30), color=(120, 120, 120)).save(buf, format="JPEG")
    return ContentFile(buf.getvalue(), name=name)


# Keep the suite off object storage. Without this, running with production
# settings would upload test fixtures into the live Spaces bucket.
local_media = override_settings(
    STORAGES={
        "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
        "staticfiles": {
            "BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"
        },
    },
    MEDIA_ROOT=tempfile.mkdtemp(prefix="kms-test-media-"),
    MEDIA_URL="/media/",
)


@local_media
class SignedUrlTtlTests(TestCase):
    def test_defaults_to_fifteen_minutes(self):
        self.assertEqual(DEFAULT_SIGNED_URL_TTL, 900)

    @override_settings(MEDIA_SIGNED_URL_TTL=300)
    def test_reads_setting(self):
        self.assertEqual(signed_url_ttl(), 300)

    @override_settings(MEDIA_SIGNED_URL_TTL="not-a-number")
    def test_falls_back_when_setting_is_junk(self):
        self.assertEqual(signed_url_ttl(), DEFAULT_SIGNED_URL_TTL)

    @override_settings(MEDIA_SIGNED_URL_TTL=0)
    def test_rejects_non_positive_ttl(self):
        self.assertEqual(signed_url_ttl(), DEFAULT_SIGNED_URL_TTL)

    def test_blank_name_has_no_url(self):
        self.assertIsNone(signed_media_url(""))
        self.assertIsNone(signed_media_url(None))

    def test_local_storage_returns_plain_media_url(self):
        """Development runs on local disk, where there is nothing to sign."""
        self.assertEqual(signed_media_url("account/documents/1/ktp/x.jpg"),
                         "/media/account/documents/1/ktp/x.jpg")


@local_media
class DocumentFileAccessTests(TestCase):
    def setUp(self):
        # DRF is configured with JWT authentication only, so session login does
        # not apply; force_authenticate bypasses the auth class while still
        # exercising the permission classes.
        self.client = APIClient()
        self.doc_type = DocumentType.objects.create(code="ktp", name="KTP")

        self.applicant = CustomUser.objects.create_user(
            email="pelamar.doc@example.com",
            password="testpass123",
            role=UserRole.APPLICANT,
            full_name="Budi Santoso",
            is_active=True,
            email_verified=True,
        )
        self.profile = ApplicantProfile.objects.create(
            user=self.applicant,
            contact_phone="081234567890",
            nik="1234567890123456",
        )
        self.document = ApplicantDocument.objects.create(
            applicant_profile=self.profile,
            document_type=self.doc_type,
            file=_tiny_jpeg(),
        )

        self.other_applicant = CustomUser.objects.create_user(
            email="pelamar.lain@example.com",
            password="testpass123",
            role=UserRole.APPLICANT,
            full_name="Siti Aminah",
            is_active=True,
            email_verified=True,
        )
        self.other_profile = ApplicantProfile.objects.create(
            user=self.other_applicant,
            contact_phone="081200000000",
            nik="1234567890123457",
        )

        self.admin = CustomUser.objects.create_user(
            email="admin.doc@example.com",
            password="testpass123",
            role=UserRole.ADMIN,
            full_name="Admin Satu",
            is_active=True,
            email_verified=True,
        )

    def _admin_file_url_path(self, document=None):
        return reverse(
            "account:applicant-document-file-url",
            kwargs={
                "applicant_pk": (document or self.document).applicant_profile.user_id,
                "pk": (document or self.document).pk,
            },
        )

    def _admin_file_path(self):
        return reverse(
            "account:applicant-document-file",
            kwargs={
                "applicant_pk": self.document.applicant_profile.user_id,
                "pk": self.document.pk,
            },
        )

    # -- authentication ---------------------------------------------------

    def test_file_url_requires_authentication(self):
        response = self.client.get(self._admin_file_url_path())
        self.assertIn(response.status_code, (401, 403))

    def test_redirect_requires_authentication(self):
        response = self.client.get(self._admin_file_path())
        self.assertIn(response.status_code, (401, 403))

    def test_applicant_cannot_use_admin_route(self):
        self.client.force_authenticate(user=self.applicant)
        response = self.client.get(self._admin_file_url_path())
        self.assertIn(response.status_code, (401, 403))

    # -- admin access -----------------------------------------------------

    def test_admin_gets_url_and_ttl(self):
        self.client.force_authenticate(user=self.admin)
        response = self.client.get(self._admin_file_url_path())
        self.assertEqual(response.status_code, 200)
        self.assertIn("url", response.json())
        self.assertEqual(response.json()["expires_in"], signed_url_ttl())

    def test_admin_url_response_is_not_cacheable(self):
        self.client.force_authenticate(user=self.admin)
        response = self.client.get(self._admin_file_url_path())
        self.assertEqual(response["Cache-Control"], "private, no-store")

    def test_redirect_points_at_the_file_and_is_not_cacheable(self):
        self.client.force_authenticate(user=self.admin)
        response = self.client.get(self._admin_file_path())
        self.assertEqual(response.status_code, 302)
        self.assertIn("ktp", response["Location"])
        self.assertEqual(response["Cache-Control"], "private, no-store")

    def test_missing_file_is_reported_as_not_found(self):
        empty = ApplicantDocument.objects.create(
            applicant_profile=self.other_profile,
            document_type=DocumentType.objects.create(code="ijasah", name="Ijasah"),
            file="",
        )
        self.client.force_authenticate(user=self.admin)
        response = self.client.get(self._admin_file_url_path(empty))
        self.assertEqual(response.status_code, 404)

    # -- applicant self-service -------------------------------------------

    def test_applicant_can_reach_own_document(self):
        self.client.force_authenticate(user=self.applicant)
        path = reverse(
            "account:applicant-me-documents-file-url",
            kwargs={"pk": self.document.pk},
        )
        response = self.client.get(path)
        self.assertEqual(response.status_code, 200)
        self.assertIn("url", response.json())

    def test_applicant_cannot_reach_another_applicants_document(self):
        other_document = ApplicantDocument.objects.create(
            applicant_profile=self.other_profile,
            document_type=self.doc_type,
            file=_tiny_jpeg("ktp-lain.jpg"),
        )
        self.client.force_authenticate(user=self.applicant)
        path = reverse(
            "account:applicant-me-documents-file-url",
            kwargs={"pk": other_document.pk},
        )
        response = self.client.get(path)
        self.assertEqual(response.status_code, 404)

    # -- serializer -------------------------------------------------------

    def test_serializer_exposes_a_stable_access_url(self):
        self.client.force_authenticate(user=self.admin)
        list_path = reverse(
            "account:applicant-documents",
            kwargs={"applicant_pk": self.applicant.pk},
        )
        response = self.client.get(list_path)
        self.assertEqual(response.status_code, 200)

        payload = response.json()
        rows = payload if isinstance(payload, list) else payload.get("results", [])
        row = next(r for r in rows if r["id"] == self.document.pk)

        self.assertIn("file_access_url", row)
        self.assertTrue(row["file_access_url"].endswith(self._admin_file_url_path()))

    def test_serializer_exposes_a_browser_view_url(self):
        """The field the frontend links to, distinct from the JSON one."""
        self.client.force_authenticate(user=self.admin)
        list_path = reverse(
            "account:applicant-documents",
            kwargs={"applicant_pk": self.applicant.pk},
        )
        response = self.client.get(list_path)
        self.assertEqual(response.status_code, 200)

        payload = response.json()
        rows = payload if isinstance(payload, list) else payload.get("results", [])
        row = next(r for r in rows if r["id"] == self.document.pk)

        self.assertTrue(row["file_view_url"].endswith(self._admin_file_path()))
        self.assertNotEqual(row["file_view_url"], row["file_access_url"])


class PrivateDocumentStorageTests(TestCase):
    """
    How new document uploads are stored.

    Asserted on the configuration rather than a real upload, because the point
    of these settings is what they tell Spaces to do; constructing the storage
    does not open a connection.
    """

    def tearDown(self):
        reset_storage_cache()

    def test_local_disk_uses_the_default_storage(self):
        """Development and this suite must not be pushed onto S3."""
        self.assertFalse(remote_storage_configured())
        self.assertIs(private_document_storage(), default_storage)

    @override_settings(AWS_STORAGE_BUCKET_NAME="kms-data")
    def test_new_uploads_are_private(self):
        reset_storage_cache()
        self.assertEqual(private_document_storage().default_acl, "private")

    @override_settings(AWS_STORAGE_BUCKET_NAME="kms-data")
    def test_urls_are_signed_against_the_origin(self):
        """
        A signature is computed over the request host, and the CDN would keep
        serving a cached copy of an object that has since become private.
        """
        reset_storage_cache()
        storage = private_document_storage()
        self.assertIsNone(storage.custom_domain)
        self.assertTrue(storage.querystring_auth)
        self.assertEqual(storage.querystring_expire, signed_url_ttl())

    @override_settings(AWS_STORAGE_BUCKET_NAME="kms-data")
    def test_documents_are_not_cached_at_the_edge(self):
        """Overrides the global max-age=86400, which is wrong for private files."""
        reset_storage_cache()
        cache_control = private_document_storage().object_parameters["CacheControl"]
        self.assertIn("no-store", cache_control)
        self.assertNotIn("86400", cache_control)

    @override_settings(AWS_STORAGE_BUCKET_NAME="kms-data")
    def test_storage_is_reused(self):
        reset_storage_cache()
        self.assertIs(private_document_storage(), private_document_storage())


@local_media
class DocumentLinkTests(TestCase):
    """
    Nothing user-facing may hand out the stored object URL.

    Every link here has to survive the objects becoming private, which means
    pointing at the file/ endpoint rather than at Spaces.
    """

    def setUp(self):
        self.factory = RequestFactory()
        self.paspor = DocumentType.objects.create(code="paspor", name="Paspor")

        self.applicant = CustomUser.objects.create_user(
            email="pelamar.link@example.com",
            password="testpass123",
            role=UserRole.APPLICANT,
            full_name="Dewi Lestari",
            is_active=True,
            email_verified=True,
        )
        self.profile = ApplicantProfile.objects.create(
            user=self.applicant,
            contact_phone="081211112222",
            nik="3210987654321098",
        )
        self.document = ApplicantDocument.objects.create(
            applicant_profile=self.profile,
            document_type=self.paspor,
            file=_tiny_jpeg("paspor.jpg"),
        )

    def _request(self):
        return self.factory.get("/", secure=True)

    # -- the shared URL builder -------------------------------------------

    def test_endpoint_uses_the_user_id_not_the_profile_pk(self):
        """
        The nested route filters on applicant_profile__user_id, so passing the
        profile pk would silently build a URL that 404s.
        """
        url = document_view_endpoint(self.profile.user_id, self.document.pk)
        self.assertEqual(
            url,
            reverse(
                "account:applicant-document-file",
                kwargs={
                    "applicant_pk": self.applicant.pk,
                    "pk": self.document.pk,
                },
            ),
        )

    def test_endpoint_is_absolute_when_given_a_request(self):
        url = document_view_endpoint(
            self.profile.user_id, self.document.pk, request=self._request()
        )
        self.assertTrue(url.startswith("http"))
        self.assertTrue(url.endswith("/file/"))

    def test_endpoint_is_none_without_ids(self):
        self.assertIsNone(document_view_endpoint(None, self.document.pk))
        self.assertIsNone(document_view_endpoint(self.profile.user_id, None))

    # -- passport link on job applications --------------------------------

    def test_passport_url_is_the_endpoint_not_the_object_url(self):
        from main.models import JobApplication
        from main.serializers import JobApplicationSerializer

        # Unsaved is enough: the field only reads obj.applicant.
        job_app = JobApplication(applicant=self.profile)
        serializer = JobApplicationSerializer(context={"request": self._request()})
        url = serializer.get_passport_file_url(job_app)

        self.assertIsNotNone(url)
        self.assertTrue(url.endswith(f"/documents/{self.document.pk}/file/"))
        self.assertNotIn("/media/", url)

    def test_passport_url_is_none_without_a_passport(self):
        from main.models import JobApplication
        from main.serializers import JobApplicationSerializer

        self.document.delete()
        job_app = JobApplication(applicant=self.profile)
        serializer = JobApplicationSerializer(context={"request": self._request()})
        self.assertIsNone(serializer.get_passport_file_url(job_app))

    # -- excel export ------------------------------------------------------

    def test_excel_links_to_the_endpoint_not_the_object_url(self):
        from openpyxl import load_workbook

        from account.services.export import generate_applicants_excel

        queryset = CustomUser.objects.filter(pk=self.applicant.pk).select_related(
            "applicant_profile__user"
        )
        workbook = load_workbook(
            generate_applicants_excel(queryset, self._request())
        )
        values = [
            str(cell.value)
            for row in workbook.active.iter_rows()
            for cell in row
            if cell.value
        ]

        expected = f"/documents/{self.document.pk}/file/"
        self.assertTrue(
            any(expected in value for value in values),
            f"No cell links to {expected}",
        )
        self.assertFalse(
            any("/media/" in value for value in values),
            "An export cell still holds a stored object URL",
        )
